"""File-type router. Each extractor returns a bounded, safe summary of the file."""
import csv
import io
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from .config import MAX_TEXT_CHARS, MAX_AUDIO_BYTES
from .groq_api import transcribe

KINDS = {
    "document": {".pdf", ".docx", ".txt", ".rtf", ".md"},
    "spreadsheet": {".xlsx", ".xls", ".csv"},
    "presentation": {".pptx"},
    "image": {".jpg", ".jpeg", ".png", ".webp", ".heic"},
    "audio": {".mp3", ".wav", ".m4a", ".ogg"},
    "video": {".mp4", ".mov", ".avi", ".mkv"},
    "code": {".py", ".js", ".ts", ".html", ".css", ".json"},
    "archive": {".zip", ".rar", ".7z"},
}


def kind_of(ext):
    for kind, exts in KINDS.items():
        if ext in exts:
            return kind
    return "other"


def _to_jpeg(data_or_path, max_side=1024):
    """Convert any supported image into a small JPEG for the vision model."""
    from PIL import Image
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
    except Exception:
        pass
    img = Image.open(data_or_path).convert("RGB")
    img.thumbnail((max_side, max_side))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=80)
    return buf.getvalue()


def _pdf(path):
    import fitz
    doc = fitz.open(path)
    text = "\n".join(doc[i].get_text() for i in range(min(5, len(doc))))
    out = {"text": text[:MAX_TEXT_CHARS], "meta": {"pages": len(doc)}}
    if len(text.strip()) < 50 and len(doc):
        png = doc[0].get_pixmap(dpi=100).tobytes("png")
        out["images"] = [("image/jpeg", _to_jpeg(io.BytesIO(png)))]
        out["meta"]["scanned"] = True
    return out


def _docx(path):
    import docx
    d = docx.Document(path)
    return {"text": "\n".join(p.text for p in d.paragraphs if p.text.strip())[:MAX_TEXT_CHARS]}


def _text(path):
    return {"text": Path(path).read_text(encoding="utf-8", errors="replace")[:MAX_TEXT_CHARS]}


def _rtf(path):
    from striprtf.striprtf import rtf_to_text
    raw = Path(path).read_text(encoding="utf-8", errors="replace")
    return {"text": rtf_to_text(raw)[:MAX_TEXT_CHARS]}


def _xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets[:3]:
        parts.append(f"Sheet: {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 6:
                break
            parts.append(" | ".join("" if c is None else str(c) for c in row))
    return {"text": "\n".join(parts)[:MAX_TEXT_CHARS], "meta": {"sheets": wb.sheetnames}}


def _xls(path):
    import xlrd
    wb = xlrd.open_workbook(path, on_demand=True)
    parts = []
    for name in wb.sheet_names()[:3]:
        sh = wb.sheet_by_name(name)
        parts.append(f"Sheet: {name}")
        for r in range(min(6, sh.nrows)):
            parts.append(" | ".join(str(c) for c in sh.row_values(r)))
    return {"text": "\n".join(parts)[:MAX_TEXT_CHARS]}


def _csv(path):
    lines = []
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        for i, row in enumerate(csv.reader(f)):
            if i >= 8:
                break
            lines.append(" | ".join(row))
    return {"text": "\n".join(lines)[:MAX_TEXT_CHARS]}


def _pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    parts = []
    for n, slide in enumerate(prs.slides, 1):
        if n > 15:
            break
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        parts.append(f"Slide {n}: " + " / ".join(texts))
    return {"text": "\n".join(parts)[:MAX_TEXT_CHARS], "meta": {"slides": len(prs.slides)}}


def _image(path):
    return {"images": [("image/jpeg", _to_jpeg(path))]}


def _audio(path):
    p = Path(path)
    if p.stat().st_size > MAX_AUDIO_BYTES:
        return {"meta": {"note": "audio too large for transcription"}}
    text = transcribe(p.name, p.read_bytes())
    if text:
        return {"text": ("Transcript: " + text)[:MAX_TEXT_CHARS]}
    return {"meta": {"note": "no transcript available"}}


def _video(path):
    """Samples up to 3 frames and the first 5 minutes of audio."""
    if not shutil.which("ffmpeg"):
        return {"meta": {"note": "ffmpeg not available"}}
    tmp = Path(tempfile.mkdtemp())
    out = {}
    try:
        subprocess.run(
            ["ffmpeg", "-loglevel", "error", "-i", str(path), "-vf", "fps=1/10,scale=640:-1",
             "-frames:v", "3", str(tmp / "f_%d.jpg")],
            timeout=60, check=False,
        )
        frames = sorted(tmp.glob("f_*.jpg"))
        if frames:
            out["images"] = [("image/jpeg", f.read_bytes()) for f in frames]
        audio = tmp / "a.mp3"
        subprocess.run(
            ["ffmpeg", "-loglevel", "error", "-i", str(path), "-t", "300", "-vn", "-ac", "1",
             "-ar", "16000", "-b:a", "32k", str(audio)],
            timeout=120, check=False,
        )
        if audio.exists() and 0 < audio.stat().st_size <= MAX_AUDIO_BYTES:
            text = transcribe("a.mp3", audio.read_bytes())
            if text:
                out["text"] = ("Transcript: " + text)[:MAX_TEXT_CHARS]
        return out
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _archive(path, ext):
    """Only reads the file listing. Contents are never extracted."""
    names = []
    if ext == ".zip":
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
    elif ext == ".7z":
        import py7zr
        with py7zr.SevenZipFile(path) as z:
            names = z.getnames()
    elif ext == ".rar":
        import rarfile
        with rarfile.RarFile(path) as z:
            names = z.namelist()
    return {"text": "Archive entries:\n" + "\n".join(names[:60]), "meta": {"entries": len(names)}}


def extract(path, ext):
    kind = kind_of(ext)
    try:
        if ext == ".pdf":
            data = _pdf(path)
        elif ext == ".docx":
            data = _docx(path)
        elif ext in {".txt", ".md"} or kind == "code":
            data = _text(path)
        elif ext == ".rtf":
            data = _rtf(path)
        elif ext == ".xlsx":
            data = _xlsx(path)
        elif ext == ".xls":
            data = _xls(path)
        elif ext == ".csv":
            data = _csv(path)
        elif ext == ".pptx":
            data = _pptx(path)
        elif kind == "image":
            data = _image(path)
        elif kind == "audio":
            data = _audio(path)
        elif kind == "video":
            data = _video(path)
        elif kind == "archive":
            data = _archive(path, ext)
        else:
            data = {}
        data["error"] = None
    except Exception as exc:
        data = {"error": f"{type(exc).__name__}: {exc}"}
    data["kind"] = kind
    return data
